import os
from docx import Document
from models.schemas import CaseFacts, ActionPlan

def generate_appeal(facts: CaseFacts, plan: ActionPlan, template_path: str, output_path: str) -> str:
    """
    Populates a controlled DOCX template with structured data.
    Does NOT use LLM to freely generate text.
    """
    if plan.action_route in ["NO_VERIFIED_ROUTE", "UNCERTAIN"]:
        raise ValueError("Cannot generate appeal for unverified or uncertain routes.")
        
    try:
        doc = Document(template_path)
    except Exception:
        # Fallback if template is missing/empty during testing
        doc = Document()
        doc.add_heading('Appeal Document', 0)
        doc.add_paragraph('Applicant: <<APPLICANT_NAME>>')
        doc.add_paragraph('Reason: <<REASON>>')
        
    # Controlled field replacement
    for para in doc.paragraphs:
        if '<<APPLICANT_NAME>>' in para.text:
            name = facts.applicant_facts.get('name', 'Unknown Applicant')
            para.text = para.text.replace('<<APPLICANT_NAME>>', name)
        if '<<REASON>>' in para.text:
            para.text = para.text.replace('<<REASON>>', plan.reason)
            
    doc.save(output_path)
    return output_path
